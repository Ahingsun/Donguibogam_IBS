"""D2. Independent expert rating: stratified sample, blinded rater workbooks, and analysis.

Design (config["d2"]):
  stratum A  passages included by the author                  rated in full
  stratum B  candidates excluded by the author                 random sample
  stratum C  candidates removed by the exclusion filter        random sample
  stratum D2 non-candidates with screening score 2             random sample
  stratum D1 non-candidates with screening score 1             random sample
  stratum E  score-0 passages in sections holding final passages   random sample
  stratum F  other score-0 passages                            not sampled (stated as an assumption)
Strata are fixed when the sample is drawn. Metrics use the known inclusion probabilities, so they
stay valid if the author's decisions change afterwards; passages the author newly includes later
are rated in an addendum so that every included passage has an expert label.

Reference label: agreement of the two raters. For the units on which they disagreed, the first author assigned a
label after reviewing both ratings (researcher decision 2026-10-27, instead of a third rater; d2/불일치_저자판정.xlsx).
These labels are not independent of the decisions under evaluation, so they enter a secondary analysis only
(with_author_labels); the primary estimates use the units on which the raters agreed.

Rating unit (second version of the rater workbooks, 2026-10-09): the unit the author adjudicates. A sampled body
passage with two or more source citations is rated unit by unit (src/statements.py), every other passage whole.
The passage stays the sampling unit, so intervals resample passages within strata with their units together.

Without the workbooks (the code as published, without the source text) the key and the answers are read from
rating_key.csv and ratings.csv, which run_d2.py analyze writes from the workbooks.
"""
import csv
import math
import os

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import statements

WRAP = Alignment(wrap_text=True, vertical="top")
FILL = PatternFill("solid", fgColor="FFF2CC")
Q_COLS = ["Q1_주된내용_배변관련", "Q2a_배변연관통증", "Q2b_빈도변화", "Q2c_형태변화", "Q2d_만성표지",
          "Q3_제외맥락", "Q4_종합판단", "Q4_확신도", "Q5_배변양상", "메모"]
STRATA = ["A", "B", "C", "D2", "D1", "E", "F"]
THIRD_OUT = "동의보감_IBS조문_제3검토위원.xlsx"          # tools/make_third_rater_file.py writes it to major_revision/
AUTHOR_LABELS = "불일치_저자판정.xlsx"   # that workbook, filled in by the first author (2026-10-27)
KEY_FILE = "평정키_연구자보관.xlsx"
KEY_CSV = "rating_key.csv"            # the key without the workbook
RATINGS_CSV = "ratings.csv"           # include/exclude answers by rating id, without the passages
RATER_COLS = ("rater_1", "rater_2")
AUTHOR_COL = "author_label"           # the first author's label for the units on which the raters disagreed


def build_frame(screen_rows, units_by_id, decisions, final_sections):
    frame = {s: [] for s in STRATA}
    for r in screen_rows:
        uid = r["unit_id"]
        u = units_by_id[uid]
        if r["candidate"]:
            if r["excluded_by_filter"]:
                frame["C"].append(uid)
            else:
                frame["A" if decisions.get(uid) else "B"].append(uid)
        elif r["score"] == 2:
            frame["D2"].append(uid)
        elif r["score"] == 1:
            frame["D1"].append(uid)
        else:
            frame["E" if u["BB"] in final_sections else "F"].append(uid)
    return frame


def draw(frame, sizes, seed, n_calibration):
    rng = np.random.default_rng(seed)
    sample, prob = [], {}
    rest = []
    for s in STRATA:
        pop = sorted(frame[s])
        k = len(pop) if sizes.get(s) == "all" else min(int(sizes.get(s, 0)), len(pop))
        idx = rng.permutation(len(pop))
        chosen = [pop[i] for i in idx[:k]]
        sample += [(uid, s) for uid in chosen]
        prob[s] = (k / len(pop)) if pop else 0.0
        if s in ("B", "C", "D2"):
            rest += [pop[i] for i in idx[k:]]
    rest = sorted(rest)
    cal = [rest[i] for i in rng.permutation(len(rest))[:n_calibration]]
    order = rng.permutation(len(sample))
    sample = [sample[i] for i in order]
    return sample, prob, cal


TRANS_COL = "원문 해석 (AI 번역·참고용)"
FILL_AID = PatternFill("solid", fgColor="DDEBF7")


def _sheet(ws, rows, units_by_id, start_id, translations=None):
    """Rater sheet. Columns are located by name, so a translation column can sit next to the text."""
    hdr = ["평정ID", "편·권", "門", "소분류", "세부", "유형", "처방명", "원문"] +         ([TRANS_COL] if translations is not None else []) + Q_COLS
    pos = {h: i + 1 for i, h in enumerate(hdr)}
    L = get_column_letter
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = WRAP
    for i, uid in enumerate(rows, start_id):
        u = units_by_id[uid]
        ws.append([i, u["book"], u["BB"], u["CC"], u["DD"], {"body": "본문", "presc": "처방"}.get(u["type"], u["type"]),
                   u["name"], u["text"]] + ([(translations or {}).get(uid, "")] if translations is not None else [])
                  + [""] * len(Q_COLS))
    widths = {"평정ID": 8, "편·권": 11, "門": 9, "소분류": 12, "세부": 9, "유형": 6, "처방명": 12,
              "원문": 45 if translations is not None else 80, TRANS_COL: 60, "메모": 30}
    for h, j in pos.items():
        ws.column_dimensions[L(j)].width = widths.get(h, 12)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            name = hdr[c.column - 1]
            if name in Q_COLS:
                c.fill = FILL
            elif name == TRANS_COL:
                c.fill = FILL_AID
    ws.freeze_panes = "B2" if translations is not None else "I2"
    n = len(rows) + 1
    col = lambda h: L(pos[h])
    yn = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True)
    y2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    q4 = DataValidation(type="list", formula1='"포함,제외"', allow_blank=True)
    cf = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    q5 = DataValidation(type="list", formula1='"설사,변비,혼합,판단불가"', allow_blank=True)
    for dv in (yn, y2, q4, cf, q5):
        ws.add_data_validation(dv)
    yn.add(f"{col(Q_COLS[0])}2:{col(Q_COLS[0])}{n}")
    y2.add(f"{col(Q_COLS[1])}2:{col(Q_COLS[4])}{n}")
    q4.add(f"{col(Q_COLS[6])}2:{col(Q_COLS[6])}{n}")
    cf.add(f"{col(Q_COLS[7])}2:{col(Q_COLS[7])}{n}")
    q5.add(f"{col(Q_COLS[8])}2:{col(Q_COLS[8])}{n}")


INSTRUCTIONS = [
    "동의보감 조문 독립 평정 — 평가자 안내",
    "",
    "목적: 각 조문이 아래 정의의 'IBS 관련 증상 조문'에 해당하는지 독립적으로 판단합니다. 다른 평가자와 상의하지 마세요.",
    "",
    "정의: 조문의 주된 임상 내용이 만성·반복성 기능성 장 증상이고, Rome IV 증상 영역(배변과 연관된 복통·불편, 배변 횟수의 변화,",
    "대변 형태의 변화)에 해당하며, 기질 질환이나 급성 질환을 가리키는 내용이 없는 조문. IBS 진단을 뜻하지 않습니다.",
    "",
    "순서: ① '보정' 시트 20건을 먼저 평정하고 연구 책임자와 기준을 맞춥니다. ② 그다음 '평정' 시트를 혼자 평정합니다.",
    "",
    "칸 설명 (노란 칸만 입력)",
    "파란 칸 '원문 해석'은 AI(Claude)가 만든 참고용 번역입니다. 판단은 원문을 기준으로 하고, 번역이 원문과 다르면 메모에 적습니다.",
    "  [?]는 번역이 불확실한 곳, [원문 확인]은 오자로 보이는 곳입니다. 원문의 '灸'는 대부분 '炙(굽다)'의 통용 표기라 '구운 것'으로 옮겼습니다.",
    "Q1  조문의 주된 임상 내용이 배변 이상 또는 배변과 관련된 복부 증상인가 — 예/아니오",
    "Q2a 배변과 연관된 복통·불편이 원문에 있는가 (예: 痛瀉, 瀉後痛減, 得後與氣則快然如衰) — Y/N",
    "Q2b 배변 횟수 변화가 원문에 있는가 (예: 泄瀉, 自利, 秘結, 大便難) — Y/N",
    "Q2c 대변 형태 변화가 원문에 있는가 (예: 溏, 鶩溏, 完穀不化, 燥結, 白沫) — Y/N",
    "Q2d 만성·반복을 뜻하는 표현이 원문에 있는가 (예: 久, 經年, 休息, 屢發, 五更) — Y/N",
    "Q3  해당하는 제외 맥락을 모두 적습니다: 혈변, 급성열병, 역병, 기생충, 임신산후, 소아, 항문질환, 전신질환, 이론맥법, 없음",
    "Q4  종합 판단: 위 정의의 IBS 관련 증상 조문으로 '포함' 또는 '제외'. 확신도 1(낮음)–3(높음)",
    "Q5  포함이면 배변 양상: 설사 / 변비 / 혼합 / 판단불가",
    "",
    "주의",
    "- 원문에 적힌 내용만으로 판단합니다. 처방 조문은 처방의 주치(治…) 문구를 기준으로 봅니다.",
    "- 조문이 여러 문장을 묶은 긴 블록이면, 블록의 주된 내용을 기준으로 판단합니다.",
    "- 파일 구조(시트 이름, 열 순서)는 바꾸지 마세요.",
]


def build_materials(sample, cal, units_by_id, out_dir, n_raters=2, translations=None):
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for k in range(n_raters):
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "안내"
        for i, t in enumerate(INSTRUCTIONS, 1):
            ws0.cell(row=i, column=1, value=t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
        ws0.column_dimensions["A"].width = 150
        _sheet(wb.create_sheet("보정"), cal, units_by_id, start_id=9001, translations=translations)
        _sheet(wb.create_sheet("평정"), [uid for uid, _ in sample], units_by_id, start_id=1, translations=translations)
        p = os.path.join(out_dir, f"평정지_평가자{chr(65 + k)}.xlsx")
        wb.save(p)
        paths.append(p)
    key = Workbook()
    ws = key.active
    ws.title = "평정키"
    ws.append(["평정ID", "unit_id", "stratum", "inclusion_probability"])
    return paths, key


def save_key(key, sample, prob, cal, out_dir):
    ws = key["평정키"]
    for i, (uid, s) in enumerate(sample, 1):
        ws.append([i, uid, s, prob[s]])
    for i, uid in enumerate(cal, 9001):
        ws.append([i, uid, "calibration", ""])
    p = os.path.join(out_dir, "평정키_연구자보관.xlsx")
    key.save(p)
    return p


# --------------------------------------------------------------------------- analysis
def read_key(folder):
    """Key rows (rating id, unit id, stratum, inclusion probability) from the key workbook, else from rating_key.csv."""
    p = os.path.join(folder, KEY_FILE)
    if os.path.exists(p):
        wb = load_workbook(p, read_only=True)
        rows = list(wb["평정키"].iter_rows(values_only=True))[1:]
        wb.close()
    else:
        with open(os.path.join(folder, KEY_CSV), encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))[1:]
    return [(int(r[0]), r[1], r[2], (float(r[3]) if r[3] not in (None, "") else None)) for r in rows]


def write_key_csv(key, folder):
    with open(os.path.join(folder, KEY_CSV), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rating_id", "unit_id", "stratum", "inclusion_probability"])
        w.writerows([rid, uid, s, "" if p is None else repr(p)] for rid, uid, s, p in key)


def write_ratings(items, raters, folder, author=None):
    """ratings.csv: the include/exclude answer of each rater for every rated unit (blank: no answer), and the first
    author's label for the units on which the raters disagreed when author holds them."""
    cols = list(RATER_COLS[:len(raters)]) + ([AUTHOR_COL] if author else [])
    with open(os.path.join(folder, RATINGS_CSV), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rating_id", *cols])
        w.writerows([rid] + [judgement(r.get(rid)) or "" for r in raters + ([author] if author else [])]
                    for rid, _, _, s, _ in items if s != "calibration")


def read_ratings(folder):
    """{column: {rating id: row}} from ratings.csv for the raters and the author's labels, rows shaped like
    read_rater's."""
    with open(os.path.join(folder, RATINGS_CSV), encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return {c: {r["rating_id"]: {JUDGEMENT_COLS[0]: r[c]} for r in rows}
            for c in RATER_COLS + (AUTHOR_COL,) if rows and c in rows[0]}


def read_rater(path):
    """Ratings of the main sheet by 검토ID as text: "12" for a passage, "12-3" for its third source unit.
    The main sheet is '검토' in the second version (2026-10-09) and '평정' before."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["검토"] if "검토" in wb.sheetnames else wb["평정"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(h) for h in rows[0]]
    return {str(r[0]).strip(): dict(zip(hdr, r)) for r in rows[1:] if r and r[0] is not None}


JUDGEMENT_COLS = ("판정(포함/제외)", "Q4_종합판단")     # second version (2026-10-09: the only answer asked), first version


def judgement(row):
    """The rater's include/exclude answer of a row, whichever version of the workbook it comes from."""
    for c in JUDGEMENT_COLS:
        v = (row or {}).get(c)
        if v not in (None, ""):
            return str(v).strip()
    return None


def rating_items(key_rows, units_by_id):
    """Rows the raters judge: (rating id, unit id, block id, stratum, prob). A sampled passage divided by source
    citation (statements.is_split) is judged unit by unit as "<id>-<k>" with unit id "<unit_id>#<k>"; each unit
    keeps its passage's stratum and inclusion probability (the passage is the sampling unit)."""
    out = []
    for rid, uid, s, p in key_rows:
        st = statements.statement_units(units_by_id[uid])
        if st:
            out += [(f"{rid}-{k}", x["unit_id"], rid, s, p) for k, x in enumerate(st, 1)]
        else:
            out.append((str(rid), uid, rid, s, p))
    return out


def cohen_kappa(a, b):
    cats = sorted(set(a) | set(b))
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def gwet_ac1(a, b):
    cats = sorted(set(a) | set(b))
    n, q = len(a), len(cats)
    if n == 0 or q < 2:
        return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    pi = {c: (a.count(c) + b.count(c)) / (2 * n) for c in cats}
    pe = sum(p * (1 - p) for p in pi.values()) / (q - 1)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def _w(x):
    """Horvitz–Thompson weight of a rated unit: the inverse of the inclusion probability of its passage."""
    return 1 / x["prob"]


INCLUDED = lambda x: x["author"] == "포함"
# the steps of the procedure, as classifications of every unit of the clinical chapters: the keyword screen (score
# at or above the threshold), the exclusion filter (screened and not filtered, i.e. adjudicated) and adjudication
STEPS = (("screen", lambda x: x["screened"] == "yes"), ("filter", lambda x: x["adjudicated"] == "yes"),
         ("final", INCLUDED))


def _precision(rows, positive=INCLUDED):
    """Stratum-weighted share of the units classed positive that the reference judges relevant."""
    pos = [x for x in rows if positive(x)]
    den = sum(_w(x) for x in pos)
    return sum(_w(x) for x in pos if x["reference"] == "포함") / den if den else float("nan")


def _recall(rows, positive=INCLUDED):
    """Stratum-weighted share of the units the reference judges relevant that are classed positive."""
    rel = [x for x in rows if x["reference"] == "포함"]
    den = sum(_w(x) for x in rel)
    return sum(_w(x) for x in rel if positive(x)) / den if den else float("nan")


def _f1(p, r):
    return 2 * p * r / (p + r) if (p + r) and not (math.isnan(p) or math.isnan(r)) else float("nan")


def _cluster_ci(rows, stat, rng, n_boot, census_fixed=False):
    """Percentile 95% interval of stat(rows), resampling passages with replacement within strata; the units of a
    divided passage stay together because they share the passage's sampling. With census_fixed, a stratum sampled in
    full (inclusion probability 1) is kept as it is, because it carries no sampling variance; this is used for
    precision and recall, which estimate quantities of the finite set of passages. Agreement statistics resample
    every stratum."""
    by = {}
    for x in rows:
        by.setdefault(x["stratum"], {}).setdefault(x["block"], []).append(x)
    vals = []
    for _ in range(n_boot):
        rs = []
        for blocks in by.values():
            ids = list(blocks)
            if census_fixed and all(x["prob"] == 1 for b in blocks.values() for x in b):
                for i in ids:
                    rs += blocks[i]
                continue
            for i in rng.integers(0, len(ids), len(ids)):
                rs += blocks[ids[i]]
        v = stat(rs)
        if not math.isnan(v):
            vals.append(v)
    return f"{np.percentile(vals, 2.5):.6f}–{np.percentile(vals, 97.5):.6f}" if vals else ""


def _scores(rows, rng, n_boot, positive=INCLUDED):
    """Weighted precision and recall (census-fixed intervals) and F1 of a classification against x["reference"],
    and the unweighted κ between the classification and the reference across the rated units."""
    p, rec = _precision(rows, positive), _recall(rows, positive)
    f1 = _f1(p, rec)
    return dict(n=len(rows), n_positive=sum(1 for x in rows if positive(x)),
                precision=round(p, 6),
                precision_ci=_cluster_ci(rows, lambda rs: _precision(rs, positive), rng, n_boot, census_fixed=True),
                recall=round(rec, 6),
                recall_ci=_cluster_ci(rows, lambda rs: _recall(rs, positive), rng, n_boot, census_fixed=True),
                f1=round(f1, 6) if not math.isnan(f1) else "",
                kappa_rated_units=round(cohen_kappa(["포함" if positive(x) else "제외" for x in rows],
                                                    [x["reference"] for x in rows]), 6))


def analyze_more(per, post_rule, seed, n_boot=2000):
    """Analyses added after the ratings came back (2026-10-26): the author's decisions against each rater on every
    unit both rated, the step of the procedure that removed the relevant units it did not include (weighted share of
    all relevant units: screening threshold, exclusion filter, adjudication), and the consensus metrics without the
    units whose author decision rests on a rule adopted after the rating files were sent (post_rule: unit ids)."""
    rng = np.random.default_rng(seed + 1)
    out = {}
    for who in ("rater_a", "rater_b"):
        out[f"vs_{who}"] = _scores([dict(x, reference=x[who]) for x in per], rng, n_boot)
    done = [x for x in per if x["reference"]]
    rel = [x for x in done if x["reference"] == "포함"]
    total = sum(_w(x) for x in rel)
    step_of = lambda x: ("screening (score below the threshold)" if x["screened"] != "yes" else
                         "exclusion filter" if x["adjudicated"] != "yes" else "adjudication")
    by = {}
    for x in rel:
        if not INCLUDED(x):
            n, w = by.get(step_of(x), (0, 0.0))
            by[step_of(x)] = (n + 1, w + _w(x))
    out["consensus_misses"] = {k: dict(units=n, weighted_share_of_relevant=round(w / total, 6))
                               for k, (n, w) in sorted(by.items())}
    kept = [x for x in done if x["unit_id"] not in post_rule]
    out["consensus_without_post_rating_rules"] = dict(_scores(kept, rng, n_boot), units_removed=len(done) - len(kept))
    return out


def with_author_labels(per, labels, seed, n_boot=2000):
    """Secondary analysis: the reference completed with the first author's labels (unit id -> 포함/제외) for the units
    on which the raters disagreed. The labels are not independent of the decisions under evaluation."""
    split = [x for x in per if not x["reference"] and labels.get(x["unit_id"]) in ("포함", "제외")]
    rows = [dict(x, reference=x["reference"] or labels.get(x["unit_id"], "")) for x in per]
    rng = np.random.default_rng(seed + 2)
    done = [x for x in rows if x["reference"]]
    out = _scores(done, rng, n_boot)
    out["recall_screen"] = round(_recall(done, STEPS[0][1]), 6)
    out.update(labels=len(split), labels_differing_from_decision=sum(1 for x in split if labels[x["unit_id"]] != x["author"]))
    return out


def analyze(items, ra, rb, adjudicator, author_included, seed, n_boot=2000, status=None):
    """items: rating_items(...). author_included: unit ids (passages or "<unit_id>#<k>") the author includes.
    status: {passage id: (screened, adjudicated)} from the current screening (keyword score at or above the
    threshold; and not removed by the exclusion filter). Returns the summary and one row per rated unit.
    Precision and recall are stratum-weighted (Horvitz–Thompson) estimates for the clinical chapters outside stratum F;
    recall is given for each step of the procedure (screen, filter, final)."""
    per = []
    for rid, uid, block, s, p in items:
        if s == "calibration":
            continue
        a, b = judgement(ra.get(rid)), judgement(rb.get(rid))
        if a not in ("포함", "제외") or b not in ("포함", "제외"):     # a blank cell: unit left out of the analysis
            continue
        ref = a if a == b else (adjudicator or {}).get(uid)
        scr, adj = (status or {}).get(statements.parent_of(uid), (False, False))
        per.append(dict(rating_id=rid, unit_id=uid, block=block, stratum=s, prob=p, rater_a=a, rater_b=b,
                        reference=ref or "", author="포함" if uid in author_included else "제외",
                        screened="yes" if scr else "", adjudicated="yes" if adj else ""))
    rng = np.random.default_rng(seed)
    kap = lambda rs: cohen_kappa([x["rater_a"] for x in rs], [x["rater_b"] for x in rs])
    ac1 = lambda rs: gwet_ac1([x["rater_a"] for x in rs], [x["rater_b"] for x in rs])
    summ = dict(unit="passage, or source unit of a passage divided by source citation",
                n_rated_by_both=len(per), n_passages=len({x["block"] for x in per}),
                kappa_q4=round(kap(per), 6), kappa_q4_ci=_cluster_ci(per, kap, rng, n_boot),
                ac1_q4=round(ac1(per), 6), ac1_q4_ci=_cluster_ci(per, ac1, rng, n_boot),
                unresolved=sum(1 for x in per if not x["reference"]))
    done = [x for x in per if x["reference"]]
    final = _scores(done, rng, n_boot)
    summ.update(precision=final["precision"], precision_ci=final["precision_ci"],
                n_author_included_rated=final["n_positive"], recall=final["recall"], recall_ci=final["recall_ci"],
                f1=final["f1"], kappa_author_vs_reference=final["kappa_rated_units"])
    summ["steps"] = {name: final if name == "final" else _scores(done, rng, n_boot, positive=f) for name, f in STEPS}
    summ["assumption"] = ("score-0 passages outside the sections holding final passages (stratum F) were not sampled "
                          "and are assumed to contain no IBS-related symptom passage")
    return summ, per
