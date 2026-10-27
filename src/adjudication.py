"""Step 3. Researcher re-adjudication workbook (build and read back).

The workbook is created once (decisions/연구자_재판정_R1.xlsx). If it already exists it is never
overwritten by run_all.py: the researcher's entries are read from it (by column name), and units that
are candidates now but have no row in it are written to an addendum workbook in outputs/.
Layout of the '판정' sheet: identification and text, the Korean reading aid (when translations are
given; AI translation, data/passage_translations_ko.csv), the researcher's input columns, then the
reference columns (screening score, keyword flags, previous decision).
tools/add_translations.py rebuilds the workbook with the translation column and carries over every
entry already made, after backing up the old file.
tools/review_layout.py rearranges the sheet for review (2026-10-15) and adds two display columns (DISPLAY_COLS),
which read_decisions leaves out.
The adjudication record (write_record, decisions/adjudication_record.csv) holds the location and decision columns of
every unit without the source text. run_all.py writes it from the workbook, and reads it when the workbook is absent
(the code as published, without the source text).
"""
import csv
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import path

DECISION_COLS = ["최종판정", "포함근거", "제외코드", "R4_배변연관통증", "R4_빈도변화", "R4_형태변화", "R4_만성표지",
                 "배변양상", "메모"]
INCLUDE_CODES = [("I-1a", "배변과 연관된 복통 (痛瀉, 瀉後痛減 등). 복통이 배변 관련·배변 횟수 변화·대변 형태 변화 중 둘 이상과 "
                          "함께 서술되면 1등급(Rome IV 핵심)이다. 설사(泄瀉)·변비(秘結)는 횟수와 형태가 함께 바뀐 것으로 본다. "
                          "불편감(脹·不寬 등)은 복통이 아니다 (2026-10-10 결정)"),
                 ("I-1b", "만성·재발 설사 계열이 조문의 주증. 복통이 적혀 있지 않으면 2등급(배변 습관 변화). 복통이 "
                          "없다고 명시하면 제외(E-8)"),
                 ("I-1c", "기능성 변비 계열이 조문의 주증. 복통이 적혀 있지 않으면 2등급(배변 습관 변화). 복통이 "
                          "없다고 명시하면 제외(E-8)"),
                 ("I-1d", "腸鳴·裏急後重 등 동반 장증상. 단독으로는 포함 근거가 아니고, 설사·변비·배변 연관 복통과 함께일 때 "
                          "더하는 근거다 (2026-10-10)"),
                 ("I-2", "변증·병기가 IBS 변증틀(脾胃虛弱·肝脾不和·脾腎陽虛·寒濕·濕熱·氣滯·大腸燥結)에 해당"),
                 ("I-3", "IBS 특이 증상구가 주증이면 門 출처가 비-GI여도 포함")]
EXCLUDE_CODES = [("E-1", "사전 지정 제외 범주(기생충·역병·霍亂·임신산후·痘疹·항문직장 기질성·급성 이질·혈변). "
                          "赤痢·血痢·赤白痢는 혈변으로 본다(2026-10-09 결정 (가)). 이질(痢疾·滯下)이 주제인 단위는 피가 섞이지 "
                          "않아도 제외한다(2026-10-10 결정 A, 근거: 동의보감 '瀉與痢不同'). 동의보감이 설사로 본 水穀痢(=飱泄)와 "
                          "'似痢非痢'는 내용으로 판단하고, 설사 조문 속 痢 언급은 설사 내용으로 판정한다. 이질이 주제인지는 "
                          "門·소분류가 아니라 그 단위의 글로 본다. 혈변은 단위 어디에 적혀 있어도 제외한다(2026-10-15)"),
                 ("E-2", "소아"),
                 ("E-3", "급성 외감 열병(傷寒·暑)에 동반된 일과성 배변 이상, 또는 급성 경과. 병(설사)이 갑자기 생겼다고 "
                          "적혔거나(暴泄·暴瀉·暴下·暴水泄 같은 급성 설사, 2026-10-26 연구자 결정: 65-2) 경과를 날수로 "
                          "말하면(期十七日當劇 등) 급성 경과다. 대변이 나오는 모양이 갑작스러운 것(其來暴速·卒然而下·一時倂下)은 "
                          "급성 경과가 아니다(2026-10-26 연구자 결정: 41-1). 병인 이름만으로는 해당하지 않고, 급하다는 뜻의 "
                          "急(裏急·急痛·宜急)도 아니다. 열(發熱·身熱)과 身痛은 외감 열병의 표지일 때만 해당한다(32-3 濕瀉身痛에 "
                          "羌活, 51의 身熱은 아님). 暑泄·濕泄처럼 병인을 이름으로 단 소분류는 글에 暑나 열 표지가 있을 "
                          "때만 E-3이다(2026-10-26 연구자 결정: 39) (예외: 寒門 太陰의 虛寒 설사는 처방이 없는 서술이나 "
                          "치법 금기도 포함한다(209-1·211-1). 少陰은 예외에 들지 않는다(217·214-8))"),
                 ("E-4", "키워드 우연 매칭: 장증상이 부수증상·감별·부작용으로만 등장하고 주증은 비-IBS. 다른 주증 뒤에 "
                          "배변 증상이 짧게(或… 등) 덧붙은 것도 여기에 든다(2026-10-25 연구자 결정: 247·253). 앞에 "
                          "원인·병기(食積·酒毒·氣滯 등)만 있고 증상은 배변 증상뿐이면 포함한다(163). 여러 갈래(부위·단계·"
                          "허실·원인)를 나열한 서술에서 한 갈래가 배변 증상 묶음이나 그 기전을 서술하면 포함하고(5-2·8-3·28-1, "
                          "2026-10-26: 241-1), 배변 증상이 비-IBS 병의 증상 목록 속 한 항목이거나 병 이름만 짝지은 목록이면 "
                          "여기에 든다(193-1·256-3)"),
                 ("E-5", "순수 이론·맥법·운기·진단원칙. 병기 문장이 그 단위에서 증상이 어떻게 나타나는지(대변 모양·때·동반 "
                          "증상)나 어떻게 생기는지(원인·기전)를 서술하면 예외로 포함하고(2026-10-10 결정 (가)), 증상을 장부·"
                          "음양·한열에 짝짓기만 하거나 감별 원칙만 말하면 제외한다(2026-10-15 연구자 결정: 12-2 小腸爲泄, 31-9 "
                          "暴瀉非陽 久瀉非陰). 치법·처방 문장은 대상 증상으로 판정한다. 운기 문장도 증상 묶음과 처방을 "
                          "서술하면 포함한다(2026-10-25: 189). 난치·불치 예후만 말하는 문장은 제외하고, 예후에 치법·처방이 "
                          "함께 있으면 치법 문장으로 판정한다(2026-10-26 연구자 결정: 67-1)"),
                 ("E-6", "배변 습관 변화나 배변 연관 복통 없음: 복통·腸鳴·팽만만 있고 설사·변비·대변 형태 변화가 없거나 배변과 "
                          "연관되지 않음 (예외: 瀉後痛減 등 배변 연관 명시). 2026-10-10 '복통 단독'에서 넓힘"),
                 ("E-7", "전신질환(부종·적취·학질 등)에 딸린 설사로 처방이 주병을 겨냥"),
                 ("E-8", "복통이 없다고 명시(腹不痛 등)한 설사·변비: Rome IV의 기능성 설사·변비이며 IBS가 아니다(2026-10-25 "
                          "연구자 결정. 근거: Rome IV 기능성 설사 정의, 진료지침의 대상은 Rome 기준 IBS뿐). 복통이 적혀 "
                          "있지 않은 것은 2등급이고, '或腹不痛'처럼 아플 때도 있으면 해당하지 않는다. 2026-10-10~10-24에는 "
                          "쓰지 않았다"),
                 ("E-9", "경고 증상: 체중 감소(肉削·消瘦·羸瘦 등)나 밤·새벽에 깨는 증상(五更·夜 등)이 함께 서술됨. 기질 질환을 "
                          "배제할 수 없다 (Rome IV 진단 과정의 경고 증상, 한의 IBS 진료지침 표 7. 2026-10-10 결정). 단위 "
                          "어디에 적혀 있어도 제외한다(五更泄·晨泄, 瘦乏 등). 약 먹는 때를 말하는 夜나 낮밤 비교는 증상이 "
                          "아니다(2026-10-15)")]
EXCLUDE_LIST = '"' + ",".join(k for k, _ in EXCLUDE_CODES) + '"'

WRAP = Alignment(wrap_text=True, vertical="top")
HEAD = Font(bold=True)
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_PREFILL = PatternFill("solid", fgColor="E2EFDA")
FILL_AID = PatternFill("solid", fgColor="DDEBF7")
TRANS_COL = "원문 해석 (AI 번역·참고용)"
AI_RATIONALE = "AI 판정 근거"
AI_CONF = "AI 확신도"
PRIORITY = "검토 우선"
REVIEW = "연구자 검토"
AI_COLS = [AI_RATIONALE, AI_CONF, PRIORITY]
# display columns of the review layout (tools/review_layout.py): an Excel formula and a static join, never read
SUMMARY_COL = "판정 요약 (자동)"
PLACE_COL = "위치 (門·소분류·세부·처방명)"
DISPLAY_COLS = (SUMMARY_COL, PLACE_COL)
REVIEWED_SHEET = "검토완료"           # rows the researcher has reviewed (tools/split_reviewed.py, 2026-10-26)
TODO_SHEET = "검토필요"               # rows still needing the researcher; called '판정' until 2026-10-26 (older files)
DECISION_SHEETS = (TODO_SHEET, "판정", REVIEWED_SHEET)
# '연구자 검토': 확인/수정 = reviewed row by row; AI수용 = high-confidence AI proposal accepted without individual review
# (researcher decision 2026-10-26), after the rule checks
ACCEPTED = "AI수용"
REVIEW_LIST = '"확인,수정,' + ACCEPTED + '"'
FILL_AI = PatternFill("solid", fgColor="EDE7F6")
FILL_FLAG = PatternFill("solid", fgColor="F8CBAD")


AI_GUIDE = [
    "",
    "AI 1차 판정 (2026-10-08, 연구자 요청)",
    "- 노란 입력 칸(최종판정~배변양상)에는 AI(Claude)가 원문에 기준(코드표)을 적용해 제안한 판정이 미리 들어 있습니다.",
    "- 보라 칸 'AI 판정 근거'와 'AI 확신도'는 그 근거입니다. AI는 이전 판정을 보지 않고 판정했습니다.",
    "- '검토 우선'(주황)은 AI 제안이 이전 판정과 다르거나 확신도가 낮은 행입니다. 이 행부터 보시면 됩니다.",
    "- 각 행을 검토해 제안에 동의하면 '연구자 검토'에 '확인'을, 칸을 고쳤으면 고친 뒤 '수정'을 고릅니다.",
    "- '연구자 검토'가 비어 있는 행은 결과 계산에서 잠정값으로만 쓰입니다. 모든 행이 검토되어야 결과가 확정(FINAL)됩니다.",
    "- AI 제안 원본은 decisions/ai_proposals_R1.csv에 따로 보관합니다.",
]

STATEMENT_GUIDE = [
    "",
    "출전 단위 판정 (2026-10-09, 연구자 결정)",
    "- 출전(《…》)이 둘 이상인 본문 조문은 출전 표기마다 나눠 한 행씩 판정합니다. 번호 '28-3'은 28번 조문의 세 번째 출전 단위입니다(unit_id 끝의 #3).",
    "- 그 출전 단위 하나만 판단합니다. 앞뒤 행과 門·소분류·세부는 뜻을 이해하는 데만 쓰고, R4_ 칸은 그 행에 실제로 적힌 것만 Y로 합니다.",
    "- 나뉜 조문에 조문 단위로 해 두신 검토와 메모는 그 조문 첫 행의 메모 칸에 옮겨 두었습니다. 출전 단위 행은 모두 새로 검토해 주세요.",
    "- 출전 단위 행의 선별점수·일치 키워드와 초록 자동 표시는 그 행의 글만으로 다시 계산한 값입니다. 이전판정은 조문 단위 값입니다.",
    "- 출전 단위 AI 1차 판정은 decisions/ai_proposals_R1_출전단위.csv에 따로 보관합니다.",
    "- 검토위원 두 분도 같은 출전 단위로 검토합니다(d2/ 검토위원 파일 v2).",
]


def _instructions(ws, n_rows, n_filter, ai=False, split=None):
    """split: (passages, divided passages) when the sheet holds source-unit rows."""
    target = (f"'판정' 시트 {n_rows}행 (조문 {split[0]}개 중 출전이 둘 이상인 본문 {split[1]}개는 출전 단위로 나눔. "
              f"키워드 점수 ≥3이고 제외 필터에 걸리지 않은 조문)") if split else \
        f"'판정' 시트 {n_rows}건 (키워드 점수 ≥3이고 제외 필터에 걸리지 않은 조문)"
    lines = [
        "동의보감 IBS 관련 증상 조문 재판정 (R1) — 작성 안내",
        "",
        f"판정 대상: {target}. '필터제외' 시트 {n_filter}건은 참고용입니다.",
        "목적: 조문마다 포함/제외를 확정하고, 포함한 조문은 Rome IV 요소가 원문에 실제로 있는지 표시합니다. 조문별 판정 근거가 이 기록에서 나옵니다.",
        "",
        "작성 방법",
        "1. 노란 칸만 입력합니다. 초록 칸은 이전 판정과 키워드 자동 표시라 참고만 하세요.",
        "   - 파란 칸 '원문 해석'은 AI(Claude)가 만든 참고용 번역입니다. 판정은 원문을 기준으로 하고, 번역이 원문과 다르면 메모에 적어 주세요.",
        "     [?]는 번역이 불확실한 곳, [원문 확인]은 오자로 보이는 곳입니다. 원문의 '灸'는 대부분 '炙(굽다)'의 통용 표기라 '구운 것'으로 옮겼습니다.",
        "2. 최종판정: '포함' 또는 '제외'. 비워 두면 이전 판정이 잠정값으로 쓰입니다.",
        "3. 포함이면 '포함근거'에 코드(I-1a 등, 여러 개면 쉼표)를 넣고, R4_ 네 칸을 Y/N으로 표시합니다.",
        "   - R4_배변연관통증: 복통(痛)이 배변과 연관되거나 설사·변비·대변 형태 변화와 함께 적혀 있으면 Y (Rome IV의 '배변과 "
        "관련되거나 배변 횟수·형태 변화와 연관된 복통'). 腹不痛, 痛 없는 脹·痞·不寬(불편감), 다른 부위의 통증, 배변 증상과 "
        "다른 구절에 있는 복통은 N",
        "   - R4_빈도변화: 배변 횟수 변화(泄瀉·自利·秘結 등)가 있으면 Y",
        "   - R4_형태변화: 대변 형태 변화(溏·鶩溏·完穀不化·燥結 등)가 있으면 Y",
        "   - R4_만성표지: 만성·반복을 뜻하는 표현(久·經年·休息·屢發 등)이 있으면 Y",
        "4. 제외이면 '제외코드'(E-2 등)만 넣으면 됩니다. R4_ 칸은 비워도 됩니다.",
        "5. 배변양상: 설사 / 변비 / 혼합 / 미분류 중 하나 (포함한 조문만).",
        "6. 메모: 판단이 애매했던 이유 등 한두 줄.",
        "",
        "주의",
        "- 자동 표시(초록)는 키워드 일치일 뿐 판단이 아닙니다. 원문을 보고 결정해 주세요.",
        "- 판정이 끝난 뒤 독립 평가자(D2) 결과를 보기 전까지 이 파일을 확정본으로 둡니다.",
        "- 파일 이름과 시트 구조는 바꾸지 마세요. 코드가 이 파일을 그대로 읽습니다.",
    ] + (AI_GUIDE if ai else []) + (STATEMENT_GUIDE if split else [])
    for i, t in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 140


def _codes(ws):
    ws.append(["코드", "정의"])
    for c in ws[1]:
        c.font = HEAD
    ws.append(["[포함 근거]", ""])
    for k, v in INCLUDE_CODES:
        ws.append([k, v])
    ws.append(["", ""])
    ws.append(["[제외 코드]", ""])
    for k, v in EXCLUDE_CODES:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 100


ID_COLS = ["번호", "unit_id", "편·권", "門", "소분류", "세부", "유형", "처방명", "원문"]
REF_COLS = ["선별점수", "일치 키워드", "자동_배변연관어", "자동_복통어", "자동_대변어", "자동_만성표지", "자동_급성열병",
            "자동_전신질환", "자동_산과", "자동_소아", "자동_이론맥", "자동_배변양상", "자동_변증", "이전판정", "이전사유"]
PREFILL_COLS = ID_COLS + REF_COLS
WIDTHS = {"번호": 5, "unit_id": 14, "편·권": 9, "門": 6, "소분류": 10, "세부": 6, "유형": 5, "처방명": 10, "원문": 45,
          TRANS_COL: 60, "선별점수": 6, "일치 키워드": 18, "이전사유": 30, "메모": 30,
          AI_RATIONALE: 40, AI_CONF: 7, PRIORITY: 12, REVIEW: 8}


def build(cands, units_by_id, screen_by_id, annotations, old_by_uid, filtered, out_path, translations=None,
          proposals=None, numbers=None, split=None):
    """proposals: {unit_id: dict with the DECISION_COLS values proposed by the AI first pass, plus
    'AI 판정 근거', 'AI 확신도', '검토 우선'}. The proposed values are written into the input columns; the
    researcher confirms or changes each row and marks it in '연구자 검토'.
    cands may hold source-unit ids "<unit_id>#<k>" (src/statements.py) with their units in units_by_id;
    numbers gives the 번호 of each row (default 1, 2, ...); split = (passages, divided passages) for the guide."""
    from openpyxl.utils import get_column_letter as L
    cols = ID_COLS + ([TRANS_COL] if translations is not None else []) +         (AI_COLS + [REVIEW] if proposals is not None else []) + DECISION_COLS + REF_COLS
    pos = {c: i + 1 for i, c in enumerate(cols)}
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "안내"
    ws = wb.create_sheet("판정")
    ws.append(cols)
    for c in ws[1]:
        c.font = HEAD
        c.alignment = WRAP
    for i, uid in enumerate(cands, 1):
        u = units_by_id[uid]
        s = screen_by_id[uid]
        a = annotations[uid]
        o = old_by_uid.get(uid, {})
        old = {"kept162": "포함", "excl177": "제외", "prefilter95": "제외(이전 필터)"}.get(o.get("old_status", ""), "신규")
        if "#" in uid and old != "신규":
            old += " (조문)"
        row = {"번호": (numbers or {}).get(uid, i), "unit_id": uid, "편·권": u["book"], "門": u["BB"], "소분류": u["CC"], "세부": u["DD"],
               "유형": {"body": "본문", "presc": "처방"}.get(u["type"], u["type"]), "처방명": u["name"], "원문": u["text"],
               TRANS_COL: (translations or {}).get(uid, ""), "선별점수": s["score"],
               "일치 키워드": s["terms"].replace("|", ", "), "자동_배변연관어": a["flag_link_terms"],
               "자동_복통어": a["flag_pain_terms"], "자동_대변어": a["flag_stool_terms"], "자동_만성표지": a["flag_chronicity"],
               "자동_급성열병": a["flag_acute_febrile"], "자동_전신질환": a["flag_systemic"], "자동_산과": a["flag_obstetric"],
               "자동_소아": a["flag_paediatric"], "자동_이론맥": a["flag_theory_or_pulse"],
               "자동_배변양상": a["auto_bowel_habit"], "자동_변증": a["auto_patterns"], "이전판정": old,
               "이전사유": o.get("old_rationale", "")}
        if proposals is not None:
            row.update({k: v for k, v in (proposals.get(uid) or {}).items() if k in DECISION_COLS + AI_COLS})
        ws.append([row.get(c, "") for c in cols])
    for c in cols:
        ws.column_dimensions[L(pos[c])].width = WIDTHS.get(c, 12 if c in DECISION_COLS else 11)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            name = cols[c.column - 1]
            if name in DECISION_COLS or name == REVIEW:
                c.fill = FILL_INPUT
            elif name == TRANS_COL:
                c.fill = FILL_AID
            elif name == PRIORITY and c.value:
                c.fill = FILL_FLAG
            elif name in AI_COLS:
                c.fill = FILL_AI
            elif name in REF_COLS[2:]:
                c.fill = FILL_PREFILL
    ws.freeze_panes = "C2"
    n = len(cands) + 1
    dv1 = DataValidation(type="list", formula1='"포함,제외"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv3 = DataValidation(type="list", formula1='"설사,변비,혼합,미분류"', allow_blank=True)
    dv4 = DataValidation(type="list", formula1=EXCLUDE_LIST, allow_blank=True)
    for dv in (dv1, dv2, dv3, dv4):
        ws.add_data_validation(dv)
    col = lambda name: L(pos[name])
    dv1.add(f"{col('최종판정')}2:{col('최종판정')}{n}")
    dv4.add(f"{col('제외코드')}2:{col('제외코드')}{n}")
    dv2.add(f"{col('R4_배변연관통증')}2:{col('R4_만성표지')}{n}")
    dv3.add(f"{col('배변양상')}2:{col('배변양상')}{n}")
    if proposals is not None:
        dv5 = DataValidation(type="list", formula1='"확인,수정"', allow_blank=True)
        ws.add_data_validation(dv5)
        dv5.add(f"{col(REVIEW)}2:{col(REVIEW)}{n}")
    ws.auto_filter.ref = f"A1:{L(len(cols))}{n}"

    wf = wb.create_sheet("필터제외")
    wf.append(["unit_id", "門", "소분류", "유형", "처방명", "선별점수", "제외 범주", "이전판정", "원문"])
    for c in wf[1]:
        c.font = HEAD
    for uid in filtered:
        u = units_by_id[uid]
        s = screen_by_id[uid]
        o = old_by_uid.get(uid, {})
        old = {"kept162": "포함", "excl177": "제외", "prefilter95": "제외(이전 필터)"}.get(o.get("old_status", ""), "신규")
        wf.append([uid, u["BB"], u["CC"], u["type"], u["name"], s["score"], s["exclusion_categories"], old, u["text"]])
    for col, w in {"A": 16, "B": 9, "C": 12, "D": 6, "E": 12, "F": 6, "G": 24, "H": 12, "I": 90}.items():
        wf.column_dimensions[col].width = w
    for row in wf.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    wf.freeze_panes = "B2"
    _codes(wb.create_sheet("코드표"))
    _instructions(ws0, len(cands), len(filtered), ai=proposals is not None, split=split)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def read_decisions(wb_path):
    """Decisions by unit_id from the 판정 sheet and, when present, the 검토완료 sheet (rows the researcher has
    reviewed, moved there by tools/split_reviewed.py). A unit found in both sheets is an error."""
    if not os.path.exists(wb_path):
        return {}
    wb = load_workbook(wb_path, read_only=True, data_only=True)
    sheets = [(name, list(wb[name].iter_rows(values_only=True))) for name in DECISION_SHEETS if name in wb.sheetnames]
    wb.close()                       # read-only workbooks keep the file open until closed
    out = {}
    for name, rows in sheets:
        hdr = [str(h) if h is not None else "" for h in rows[0]]
        for r in rows[1:]:
            d = dict(zip(hdr, r))
            if d.get("unit_id"):
                uid = str(d["unit_id"])
                if uid in out:
                    raise ValueError(f"{uid} appears in more than one sheet of {os.path.basename(wb_path)}")
                out[uid] = {k: ("" if d.get(k) is None else str(d.get(k)).strip()) for k in hdr
                            if k not in DISPLAY_COLS}
    return out


# ---------------------------------------------------------------------------- adjudication record (no source text)
RECORD_ID_COLS = ["번호", "unit_id", "편·권", "門", "소분류", "세부", "유형", "처방명"]
RECORD_DECISION_COLS = [c for c in DECISION_COLS if c != "메모"]


def _text_order(uid):
    """Order of a unit in the source: volume, first line, source unit."""
    book, rest = uid.split(":")
    line, _, k = rest.partition("#")
    return book, int(line), int(k or 0)


def write_record(dec, out_path):
    """The adjudication record: location and decision columns of every unit in the workbook, in text order,
    without the source text, its translation, the AI columns and the notes."""
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=RECORD_ID_COLS + RECORD_DECISION_COLS, extrasaction="ignore")
        w.writeheader()
        for uid in sorted(dec, key=_text_order):
            w.writerow({k: dec[uid].get(k, "") for k in RECORD_ID_COLS + RECORD_DECISION_COLS})
    return out_path


def read_record(p):
    """Decisions by unit_id from the adjudication record (write_record)."""
    with open(p, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    out = {}
    for r in rows:
        uid = (r.get("unit_id") or "").strip()
        if uid:
            if uid in out:
                raise ValueError(f"{uid} appears twice in {os.path.basename(p)}")
            out[uid] = {k: (v or "").strip() for k, v in r.items()}
    return out


def load_decisions(cfg):
    """(decisions, origin): from the researcher workbook when it exists ('workbook'), else from the adjudication
    record ('record'), else ({}, None)."""
    wb = path(cfg["adjudication"]["workbook"])
    if os.path.exists(wb):
        return read_decisions(wb), "workbook"
    rec = path(cfg["adjudication"]["record"])
    if os.path.exists(rec):
        return read_record(rec), "record"
    return {}, None
